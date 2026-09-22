from __future__ import annotations

import uuid
from datetime import UTC, datetime
from pathlib import Path
from xml.etree.ElementTree import Element, SubElement, tostring

from jsonschema import ValidationError, validate
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from services.shared.models import AnswerPayload, OutputFormat, OutputRequest, RenderedArtifact


class FormattingError(ValueError):
    pass


class FormattingService:
    def __init__(self, artifact_dir: Path):
        self.artifact_dir = artifact_dir
        self.artifact_dir.mkdir(parents=True, exist_ok=True)

    def render(
        self, payloads: list[AnswerPayload], request: OutputRequest
    ) -> tuple[object | None, RenderedArtifact | None]:
        if request.target_format == OutputFormat.CHAT:
            return None, None
        if request.target_format == OutputFormat.JSON:
            content = {"answers": [payload.model_dump(mode="json") for payload in payloads]}
            if request.schema_override:
                try:
                    validate(content, request.schema_override)
                except ValidationError as error:
                    raise FormattingError(f"JSON output failed the requested schema: {error.message}") from error
            return content, None
        if request.target_format == OutputFormat.XML:
            root = Element("ueaaResponse", generatedAt=datetime.now(UTC).isoformat())
            for payload in payloads:
                answer = SubElement(root, "answer", domain=payload.domain.value)
                SubElement(answer, "text").text = payload.answer_text
                SubElement(answer, "confidence").text = str(payload.confidence)
                citations = SubElement(answer, "citations")
                for citation in payload.citations:
                    item = SubElement(citations, "citation", documentId=citation.document_id)
                    item.text = citation.excerpt
            return tostring(root, encoding="unicode"), None
        if request.target_format == OutputFormat.EMAIL:
            tone = request.tone or ("formal" if any(p.domain.value in {"hr", "legal"} for p in payloads) else "helpful")
            recipient = request.recipient_hint or "there"
            body = "\n\n".join(payload.answer_text for payload in payloads)
            return (
                f"Subject: Requested information\n\nHello {recipient},\n\n{body}\n\n"
                f"Regards,\nEnterprise AI Agent\n\n"
                f"AI-drafted ({tone}); please review before sending.",
                None,
            )
        if request.target_format == OutputFormat.XLSX:
            return None, self._workbook(payloads)
        raise FormattingError(f"Unsupported output format: {request.target_format}")

    def _workbook(self, payloads: list[AnswerPayload]) -> RenderedArtifact:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Answers"
        headers = ["Domain", "Answer", "Confidence", "Needs review", "Citation document", "Citation excerpt"]
        worksheet.append(headers)
        fill = PatternFill("solid", fgColor="003B5C")
        for cell in worksheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = fill
        for payload in payloads:
            citations = payload.citations or [None]
            for citation in citations:
                worksheet.append(
                    [
                        payload.domain.value,
                        payload.answer_text,
                        payload.confidence,
                        payload.needs_human_review,
                        citation.document_title if citation else "",
                        citation.excerpt if citation else "",
                    ]
                )
        worksheet.freeze_panes = "A2"
        for column, width in {"A": 14, "B": 60, "C": 12, "D": 14, "E": 30, "F": 70}.items():
            worksheet.column_dimensions[column].width = width
        filename = f"ueaa-export-{uuid.uuid4().hex[:12]}.xlsx"
        workbook.save(self.artifact_dir / filename)
        return RenderedArtifact(
            filename=filename,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            download_url=f"/artifacts/{filename}",
        )
